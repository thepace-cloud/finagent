"""组装网报平台提交包（对照平台表单逐栏产出，2026-10-08 硬底线日）

用法: python -X utf8 build_submit_pack.py
输出全部进入 docs/提交材料/：
  核心功能截图1/2/3-*.jpg     ← 平台"作品简介-核心功能截图"三栏（jpg≤1MB）
  作品简介-文字简介.pdf       ← 平台"文字简介"栏（PDF，汉字≤400）
  软件设计文档.pdf            ← 参赛作品文档 + AI使用声明 合并（平台仅收 PDF）
  源代码.zip                  ← 完整源码（zip 根目录直放，无媒体/无密钥/无模型）
  安装包-网页版扫码直达.zip   ← Web 作品无 APK：使用说明 + 公网二维码
依赖: docx2pdf(Word COM) pypdf qrcode pillow python-docx
"""
import os
import re
import zipfile
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
OUT = os.path.join(HERE, "提交材料")
FIG = os.path.join(HERE, "图")

# ── ① 核心功能截图 ×3 → JPG ≤1MB ────────────────────────────
SHOTS = [
    ("fig1-主对话.png", "核心功能截图1-主对话与引用.jpg"),
    ("fig3-多产品对比表.png", "核心功能截图2-多产品对比表.jpg"),
    ("fig7-记忆面板.png", "核心功能截图3-记忆面板.jpg"),
]


def make_shots():
    for src, dst in SHOTS:
        img = Image.open(os.path.join(FIG, src)).convert("RGB")
        out = os.path.join(OUT, dst)
        for q in (88, 80, 70):
            img.save(out, "JPEG", quality=q)
            if os.path.getsize(out) <= 1024 * 1024:
                break
        print(f"截图 {dst} {os.path.getsize(out)//1024}KB")


# ── ② 文字简介（汉字≤400）→ PDF ─────────────────────────────
INTRO = (
    "“贷”你明白 FinAgent 是一款面向贷款小白的可溯源 AI 问答助手。它读取三份真实贷款"
    "产品说明书，回答只基于检索到的原文生成，每个论断附[文档 页码]级引用，点击即可查看原文，"
    "用户可当场核对，从根上防治大模型“一本正经地编造”。核心功能：①贷款适配测评——答"
    "四道选择题，规则代码按说明书条款确定性打分（零大模型调用），毫秒级给出推荐产品、逐条挂"
    "引用的推荐理由与月供估算，三款均不适合时诚实全拒并逐款说明原因；②多产品对比——每款产"
    "品独立检索、逐维度对照表呈现，未检出的字段如实标注“未检出”，绝不编造；③长期记"
    "忆——面板可编辑用户画像，改“首套/二套”后同一问题的利率口径随之变化；④出口护栏"
    "——引用校验不过自动带反馈重试，仍不过则明确拒答。作品为响应式 Web 应用，无需安装注册，"
    "手机浏览器访问 http://47.95.222.61:8000 即用；运营看板实时展示请求量、三路径占比与护栏"
    "三档计数。部署采用 Docker 三容器架构，迭代全程 56 次 git 提交，自建 18 题评估集验证质量。"
)


def cjk_count(s):
    return len(re.findall(r"[一-鿿]", s))


def make_intro_pdf():
    assert cjk_count(INTRO) <= 400, f"简介超字数: {cjk_count(INTRO)}"
    print(f"简介汉字数: {cjk_count(INTRO)}")
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("“贷”你明白 FinAgent——作品简介")
    r.font.name = "黑体"
    r.font.size = Pt(16)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    body = doc.add_paragraph()
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.first_line_indent = Pt(24)
    r2 = body.add_run(INTRO)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    tmp = os.path.join(OUT, "_intro_tmp.docx")
    doc.save(tmp)
    from docx2pdf import convert
    convert(tmp, os.path.join(OUT, "作品简介-文字简介.pdf"))
    os.remove(tmp)
    print("作品简介-文字简介.pdf 完成")


# ── ③ 软件设计文档 PDF（正文 + AI 声明合并） ───────────────
def make_design_pdf():
    from docx2pdf import convert
    from pypdf import PdfWriter
    parts = []
    for name in ("参赛作品文档", "AI使用声明"):
        src = os.path.join(OUT, f"{name}.docx")   # 提交材料内副本，避开 Word 占用源文件
        tmp = os.path.join(OUT, f"_{name}.pdf")
        convert(src, tmp)
        parts.append(tmp)
    w = PdfWriter()
    for f in parts:
        w.append(f)
    dest = os.path.join(OUT, "软件设计文档.pdf")
    with open(dest, "wb") as fh:
        w.write(fh)
    for f in parts:
        os.remove(f)
    from pypdf import PdfReader
    print(f"软件设计文档.pdf 完成：{os.path.getsize(dest)//1024}KB，"
          f"{len(PdfReader(dest).pages)} 页（正文＋AI 使用声明附录）")


# ── ④ 源代码.zip（根目录直放；无媒体/无密钥/无模型/无索引产物） ──
SRC_FILES = [
    "README.md", "requirements.txt", "Dockerfile", "docker-compose.yml",
    "config.py", "llm.py", "tools.py", "react.py",
    "v3_hybrid.py", "v3_eval.py", "v3_golden.json", "v3_题单.md",
    "v3_eval_results.txt",
    "v4_db.py", "v4_store.py", "v4_memory.py", "v4_mcp_server.py",
    "v5_app.py", "v5_router.py", "v5_validator.py", "v5_retry.py",
    "v5_compare.py", "v5_assess.py", "v5_ratelimit.py", "v5_qexpand.py",
    "v5_observability.py", "v5_analysis.py", "v5_m4_traffic.py",
    "build_eval.py",
    "probe_cmp_result.json", "probe_m1_result.json", "probe_m2_result.json",
    "probe_m3_result.json", "probe_m35_result.json", "probe_m4b_result.json",
]
SRC_DIRS = ["corpus", "static"]
SRC_NOTE = """《源代码说明——原创部分与第三方部分》

一、原创部分（本团队编写的全部业务代码）
1. 应用与主链路（研发负责人独立完成，可解释每行原理）：
   v5_app.py（FastAPI 入口/SSE 流式）、v5_router.py（任务分流）、
   v3_hybrid.py（BM25+向量混合检索 RRF 融合）、v5_validator.py（引用出口校验护栏）、
   v5_retry.py（带反馈重试）、react.py（工具调用循环）、llm.py、tools.py、config.py
2. 冲刺期新增模块（AI 辅助编写；设计经本人拍板、原理经本人审读、验收由本人执行，
   详见《软件设计文档》附录《AI 使用声明》）：
   v5_compare.py（多产品对比路径）、v5_assess.py（适配测评确定性打分）、
   v5_ratelimit.py（每 IP 滑动窗口限流）、v5_qexpand.py（同义扩展实验，评测证伪后已回滚留档）、
   v5_observability.py（/metrics 看板数据）、v5_analysis.py、v5_m4_traffic.py
3. 前端（单文件、零框架）：static/index.html、static/stats.html
4. 评测体系：v3_eval.py、v3_golden.json、v3_题单.md、build_eval.py、v3_eval_results.txt
5. 语料：corpus/（三份真实贷款产品说明书的整理文本）
6. 存储层：v4_db.py、v4_store.py、v4_memory.py（会话/画像/滚动摘要）

二、第三方部分（开源依赖与外部服务，非本团队代码，已在设计文档中标明）
- FastAPI / Uvicorn（Web 框架，requirements.txt）
- sentence-transformers 及预训练向量模型（models/ 因体积未打包，运行时自动下载）
- Redis、MySQL（Docker 官方镜像）
- DeepSeek API（大模型生成服务，付费；密钥经环境变量注入，仓库不含任何密钥）
- Docker / Docker Compose（部署）

三、运行方式
docker compose up -d --build app（首次启动自动下载向量模型并建立索引）；
或在公网直接体验：http://47.95.222.61:8000（无需安装、无需注册登录、无用户名密码）
"""


def make_source_zip():
    dest = os.path.join(OUT, "源代码.zip")
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
        for f in SRC_FILES:
            p = os.path.join(REPO, f)
            if os.path.exists(p):
                z.write(p, f)
            else:
                print(f"  [跳过] {f} 不存在")
        for d in SRC_DIRS:
            base = os.path.join(REPO, d)
            for root, _, files in os.walk(base):
                if "__pycache__" in root:
                    continue
                for f in files:
                    p = os.path.join(root, f)
                    z.write(p, os.path.relpath(p, REPO))
        z.writestr("源码说明-原创与第三方.txt", SRC_NOTE)
    names = zipfile.ZipFile(dest).namelist()
    assert not any(n.startswith((".env", "models/")) for n in names)
    assert not any(n.endswith((".mp4", ".mp3", ".wav", ".log")) for n in names)
    print(f"源代码.zip 完成：{len(names)} 个文件，"
          f"{os.path.getsize(dest)//1024}KB，根目录直放")


# ── ⑤ 安装包 zip（Web 作品：使用说明 + 公网二维码） ─────────
USE_NOTE = """“贷”你明白 FinAgent——跨平台响应式 Web 应用（浏览器即用）

【访问方式】无需安装、无需注册登录、无用户名密码：
手机或电脑浏览器直接打开 http://47.95.222.61:8000
或扫描压缩包内 二维码-公网直达.png

【说明】
- 本作品为 Web 应用，不存在 Android/iOS 安装包（APK/IPA），
  按跨平台作品的网页形态提供扫码直达方式。
- 页面顶部角标实时显示后端 Redis/MySQL 状态；
  运营看板：http://47.95.222.61:8000/stats（聊天页右上角进入）。
- 评审可体验全部功能：问答、适配测评、多产品对比、记忆面板、运营看板。
"""


def make_installer_zip():
    import qrcode
    img = qrcode.make("http://47.95.222.61:8000")
    qr_buf = BytesIO()
    img.save(qr_buf, "PNG")
    dest = os.path.join(OUT, "安装包-网页版扫码直达.zip")
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("使用说明.txt", USE_NOTE)
        z.writestr("二维码-公网直达.png", qr_buf.getvalue())
    print(f"安装包-网页版扫码直达.zip 完成 {os.path.getsize(dest)//1024}KB")


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    make_shots()
    make_intro_pdf()
    make_design_pdf()
    make_source_zip()
    make_installer_zip()
    print("\n=== 提交材料目录 ===")
    for f in sorted(os.listdir(OUT)):
        print(f"  {f}  {os.path.getsize(os.path.join(OUT, f))//1024}KB")
